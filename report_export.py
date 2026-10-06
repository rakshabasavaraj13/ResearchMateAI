
from io import BytesIO
from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer
)
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.enums import TA_CENTER
from xml.sax.saxutils import escape


# ============================================================
# DOCX
# ============================================================

def create_docx(report_text):

    document = Document()

    # Normal font
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    for line in report_text.splitlines():

        line = line.strip()

        if not line:
            continue

        if line.startswith("# "):

            heading = document.add_heading(
                line[2:].strip(),
                level=0
            )

        elif line.startswith("## "):

            document.add_heading(
                line[3:].strip(),
                level=1
            )

        elif line.startswith("### "):

            document.add_heading(
                line[4:].strip(),
                level=2
            )

        else:

            document.add_paragraph(line)

    output = BytesIO()

    document.save(output)

    output.seek(0)

    return output.getvalue()


# ============================================================
# PDF
# ============================================================

def create_pdf(report_text):

    output = BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=50,
        leftMargin=50,
        topMargin=50,
        bottomMargin=50
    )

    styles = getSampleStyleSheet()

    title_style = styles["Title"]
    title_style.fontName = "Times-Roman"
    title_style.alignment = TA_CENTER

    heading_style = styles["Heading2"]
    heading_style.fontName = "Times-Roman"

    body_style = styles["BodyText"]
    body_style.fontName = "Times-Roman"
    body_style.fontSize = 11
    body_style.leading = 16

    story = []

    for line in report_text.splitlines():

        line = line.strip()

        if not line:
            story.append(Spacer(1, 8))
            continue

        safe_line = escape(line)

        if line.startswith("# "):

            story.append(
                Paragraph(
                    escape(line[2:]),
                    title_style
                )
            )

        elif line.startswith("## "):

            story.append(
                Paragraph(
                    escape(line[3:]),
                    heading_style
                )
            )

        elif line.startswith("### "):

            story.append(
                Paragraph(
                    escape(line[4:]),
                    heading_style
                )
            )

        else:

            story.append(
                Paragraph(
                    safe_line,
                    body_style
                )
            )

    document.build(story)

    output.seek(0)

    return output.getvalue()


print("✅ report_export.py created")
print("✅ PDF generator ready")
print("✅ DOCX generator ready")
